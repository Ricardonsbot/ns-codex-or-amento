"""Carrega BU -> Torre -> Sub Torre -> Empresa no Postgres, do tbl_dbEmpresas.

Fonte: FPA_DW/Datalake_readFiles/tbl_dbEmpresas.xlsx (177 empresas).

ISTO DESTRAVA A FASE 1. O BANCO-DE-DADOS.md dizia, no passo 2, que "uma coisa
não sai dos arquivos: a BU", e que o de/para Torre -> BU precisaria ser escrito
à mão por alguém do negócio. Ele existe — por empresa, nesta planilha.

TRÊS DECISÕES DE TRADUÇÃO, porque são de leitura e não de dado:

1. VALE BU2025/Torre2025, NÃO BU/Torre. As duas duplas existem e DISCORDAM: a
   antiga tem 12 valores de BU com granularidade de torre (TMS, Fintech, V&RM),
   a de 2025 tem 7 e é a hierarquia de verdade (BU EMBARCADOR, BU PSL,
   Corporate, Projetos...). Usar a errada troca o nível inteiro da árvore.

2. "RemoverDoFiltro" NÃO É SUB TORRE. Aparece 17 vezes em Torre2025_Sub_STO e é
   marcador de relatório, não nome. Vira sub torre nula — que o schema já
   permite, porque nem toda empresa pendura numa sub torre.

3. CNPJ NÃO É CHAVE. Só 59 das 177 têm. A chave natural é ID_Empresa (EMP0001…),
   que está em todas as 177 e é única. O CNPJ entra como atributo quando existe.

IDEMPOTENTE. Empresa que sumiu da origem vira INATIVA, nunca apagada: lançamento
aponta para ela, e apagar quebraria a referência e a história de um orçamento.

Uso:
    DATABASE_URL=postgresql://user:senha@localhost:5433/orcamento_dev \\
    FPA_DW_DIR="C:/.../FPA_DW/Datalake_readFiles" \\
        python ferramentas/carrega_organizacao_no_banco.py [--dry-run]
"""
import os
import re
import sys
import unicodedata

import psycopg2
import psycopg2.extras
from openpyxl import load_workbook

ARQUIVO = "tbl_dbEmpresas.xlsx"
ABA = "Empresas"

# Valores que são marcador de relatório, não nome de nada.
NAO_E_NOME = {"removerdofiltro", "na", "nda", "semclass", "-", ""}


def limpo(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "").replace("\xa0", " ")).strip()


def codigo_de(nome):
    """Código estável a partir do nome: sem acento, maiúsculo, com hífen.
    'BU EMBARCADOR' -> 'BU-EMBARCADOR'. Precisa ser estável porque é a chave
    natural de bu e torre — se mudar, a carga seguinte cria duplicata em vez de
    atualizar."""
    sem_acento = "".join(c for c in unicodedata.normalize("NFD", nome)
                         if unicodedata.category(c) != "Mn")
    return re.sub(r"-+", "-", re.sub(r"[^A-Za-z0-9]+", "-", sem_acento).strip("-")).upper()


def ler_empresas(pasta):
    wb = load_workbook(os.path.join(pasta, ARQUIVO), read_only=True, data_only=True)
    ws = wb[ABA]
    it = ws.iter_rows(values_only=True)
    cab = [limpo(c) for c in next(it)]
    linhas = [{cab[i]: limpo(r[i]) for i in range(min(len(cab), len(r)))}
              for r in it if any(v is not None for v in r)]
    wb.close()
    return linhas


def main():
    dry = "--dry-run" in sys.argv
    url = os.environ.get("DATABASE_URL")
    pasta = os.environ.get("FPA_DW_DIR")
    if not url:
        sys.exit("DATABASE_URL ausente.")
    if not pasta:
        sys.exit("FPA_DW_DIR ausente — aponte para a pasta Datalake_readFiles.")
    if not os.path.isfile(os.path.join(pasta, ARQUIVO)):
        sys.exit(f"{ARQUIVO} não encontrado em {pasta}")

    linhas = ler_empresas(pasta)
    print(f"{ARQUIVO}: {len(linhas)} empresas")

    sem_hierarquia = [r for r in linhas if not r.get("BU2025") or not r.get("Torre2025")]
    if sem_hierarquia:
        # empresa.torre_id é NOT NULL: sem torre a linha não tem onde entrar.
        print(f"  ⚠ {len(sem_hierarquia)} sem BU2025/Torre2025 — ficam de fora: "
              f"{[r.get('Company') for r in sem_hierarquia][:5]}")
        linhas = [r for r in linhas if r.get("BU2025") and r.get("Torre2025")]

    # ── Grafias diferentes do mesmo nome ────────────────────────────────────
    # "BU EMBARCADOR" (83 empresas) e "BU Embarcador" (1) são a mesma BU. Como
    # codigo_de() ignora caixa, as duas dão o mesmo código e o upsert juntaria
    # sozinho — mas o NOME gravado seria o da última linha lida, ou seja,
    # decidido por ordenação. Aqui a grafia canônica é a MAIS FREQUENTE, que é
    # escolha defensável, e a divergência fica registrada.
    def canonizar(valores):
        contagem = {}
        for v in valores:
            contagem.setdefault(codigo_de(v), {}).setdefault(v, 0)
            contagem[codigo_de(v)][v] += 1
        canon, avisos = {}, []
        for cod, grafias in contagem.items():
            vencedora = max(grafias.items(), key=lambda kv: (kv[1], kv[0]))[0]
            for g in grafias:
                canon[g] = vencedora
            if len(grafias) > 1:
                avisos.append((vencedora, {g: n for g, n in grafias.items() if g != vencedora}))
        return canon, avisos

    canon_bu, avisos_bu = canonizar([r["BU2025"] for r in linhas])
    canon_torre, avisos_torre = canonizar([r["Torre2025"] for r in linhas])
    for rotulo, avisos in (("BU", avisos_bu), ("Torre", avisos_torre)):
        for vencedora, outras in avisos:
            print(f"  ⚠ {rotulo} com mais de uma grafia — unificada como {vencedora!r}, "
                  f"absorvendo {outras}")
    for r in linhas:
        r["BU2025"] = canon_bu[r["BU2025"]]
        r["Torre2025"] = canon_torre[r["Torre2025"]]

    bus = sorted({r["BU2025"] for r in linhas})
    torres = sorted({(r["BU2025"], r["Torre2025"]) for r in linhas})

    # Nomes que NÃO colidem no código mas parecem a mesma coisa ("SW Embarcador"
    # e "Torre SW Embarcador"). Não junta sozinho: tirar o prefixo "Torre" é
    # palpite, e juntar torre errada move empresa de lugar no consolidado.
    def raiz(nome):
        return re.sub(r"^TORRE-", "", codigo_de(nome))

    porRaiz = {}
    for _, t in torres:
        porRaiz.setdefault(raiz(t), []).append(t)
    for r, nomes in porRaiz.items():
        if len(nomes) > 1:
            print(f"  ⚠ torres parecidas, NÃO unificadas (decisão sua): {nomes}")
    subs = sorted({(r["Torre2025"], r["Torre2025_Sub_STO"]) for r in linhas
                   if r.get("Torre2025_Sub_STO", "").lower() not in NAO_E_NOME})
    ativas = sum(1 for r in linhas if r.get("AtivoYN", "").upper() == "Y")

    print(f"  {len(bus)} BUs | {len(torres)} torres | {len(subs)} sub torres | "
          f"{len(linhas)} empresas ({ativas} ativas)")
    for bu in bus:
        ts = [t for b, t in torres if b == bu]
        print(f"     {bu}: {len(ts)} torre(s) — {', '.join(ts[:4])}{'…' if len(ts) > 4 else ''}")

    # Torre com o mesmo nome em BUs diferentes quebraria o UNIQUE de torre.codigo,
    # e significaria que o nome não identifica a torre. Melhor saber agora.
    por_nome = {}
    for bu, t in torres:
        por_nome.setdefault(t, []).append(bu)
    ambiguas = {t: bs for t, bs in por_nome.items() if len(bs) > 1}
    if ambiguas:
        print(f"  ⚠ torre com o mesmo nome em mais de uma BU: {ambiguas}")
        sys.exit("Nome de torre não identifica a torre — decida o código antes de carregar.")

    if dry:
        print("\n(dry-run) nada gravado.")
        return

    conn = psycopg2.connect(url)
    try:
        with conn, conn.cursor() as cur:
            cur.execute("SET search_path TO cadastro, public")

            for nome in bus:
                cur.execute("""
                    INSERT INTO cadastro.bu (codigo, nome) VALUES (%s, %s)
                    ON CONFLICT (codigo) DO UPDATE SET nome = EXCLUDED.nome, ativo = true
                """, (codigo_de(nome), nome))
            cur.execute("SELECT nome, id FROM cadastro.bu")
            id_bu = dict(cur.fetchall())

            for bu, nome in torres:
                cur.execute("""
                    INSERT INTO cadastro.torre (bu_id, codigo, nome) VALUES (%s, %s, %s)
                    ON CONFLICT (codigo) DO UPDATE
                       SET bu_id = EXCLUDED.bu_id, nome = EXCLUDED.nome, ativo = true
                """, (id_bu[bu], codigo_de(nome), nome))
            cur.execute("SELECT nome, id FROM cadastro.torre")
            id_torre = dict(cur.fetchall())

            for torre, nome in subs:
                cur.execute("""
                    INSERT INTO cadastro.sub_torre (torre_id, nome) VALUES (%s, %s)
                    ON CONFLICT (torre_id, nome) DO UPDATE SET ativo = true
                """, (id_torre[torre], nome))
            cur.execute("SELECT t.nome, s.nome, s.id FROM cadastro.sub_torre s "
                        "JOIN cadastro.torre t ON t.id = s.torre_id")
            id_sub = {(t, s): i for t, s, i in cur.fetchall()}

            registros = []
            for r in linhas:
                sub = r.get("Torre2025_Sub_STO", "")
                sub_id = id_sub.get((r["Torre2025"], sub)) if sub.lower() not in NAO_E_NOME else None
                registros.append((
                    id_torre[r["Torre2025"]],
                    sub_id,
                    r["ID_Empresa"],
                    r.get("Company") or r["ID_Empresa"],
                    r.get("CNPJ") or None,
                    r.get("AtivoYN", "").upper() == "Y",
                ))

            psycopg2.extras.execute_batch(cur, """
                INSERT INTO cadastro.empresa (torre_id, sub_torre_id, codigo, nome, cnpj, ativo)
                VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (codigo) DO UPDATE
                   SET torre_id = EXCLUDED.torre_id, sub_torre_id = EXCLUDED.sub_torre_id,
                       nome = EXCLUDED.nome, cnpj = EXCLUDED.cnpj, ativo = EXCLUDED.ativo
            """, registros, page_size=200)

            codigos = [r[2] for r in registros]
            cur.execute("""
                UPDATE cadastro.empresa SET ativo = false
                 WHERE ativo AND NOT (codigo = ANY(%s)) RETURNING codigo
            """, (codigos,))
            sumidas = [r[0] for r in cur.fetchall()]

        with conn.cursor() as cur:
            cur.execute("SELECT (SELECT count(*) FROM cadastro.bu), "
                        "(SELECT count(*) FROM cadastro.torre), "
                        "(SELECT count(*) FROM cadastro.sub_torre), "
                        "(SELECT count(*) FROM cadastro.empresa WHERE ativo)")
            nb, nt, ns, ne = cur.fetchone()
    finally:
        conn.close()

    print(f"\ncarregado: {nb} BUs, {nt} torres, {ns} sub torres, {ne} empresas ativas")
    if sumidas:
        print(f"⚠ {len(sumidas)} empresa(s) inativada(s) (sumiram da origem): "
              f"{', '.join(sumidas[:5])}{'…' if len(sumidas) > 5 else ''}")


if __name__ == "__main__":
    main()
